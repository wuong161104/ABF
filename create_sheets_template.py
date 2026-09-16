import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

def create_project_task_template():
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles & Colors
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Arial", size=14, bold=True, color="1E293B")
    subtitle_font = Font(name="Arial", size=10, italic=True, color="64748B")
    data_font = Font(name="Arial", size=10, color="000000")
    
    # Fills
    primary_header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Dark Blue
    green_header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid") # Dark Emerald
    purple_header_fill = PatternFill(start_color="5B21B6", end_color="5B21B6", fill_type="solid") # Deep Purple
    gray_header_fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid") # Slate
    accent_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # ----------------------------------------------------
    # 1. TAB: Cấu hình Mapping Topics
    # ----------------------------------------------------
    ws_map = wb.create_sheet(title="Mapping_Topics")
    ws_map.views.sheetView[0].showGridLines = True

    ws_map["A1"] = "BẢNG ÁNH XẠ TELEGRAM TOPIC -> PROJECT TAB (CHO N8N)"
    ws_map["A1"].font = title_font
    ws_map["A2"] = "Dùng bảng này để n8n tra cứu: Khi nhận tin nhắn ở Topic ID nào thì tự động ghi vào Tab Sheet tương ứng."
    ws_map["A2"].font = subtitle_font

    map_headers = ["Telegram Topic ID (message_thread_id)", "Tên Tab Sheet (Project)", "Tên Dự Án Đầy Đủ", "Leader Phụ Trách", "Ghi Chú"]
    for col_num, header in enumerate(map_headers, 1):
        cell = ws_map.cell(row=4, column=col_num)
        cell.value = header
        cell.font = header_font
        cell.fill = gray_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sample_mappings = [
        [101, "Project_Video", "Sản Xuất Video & Media", "Trần Văn A", "Topic thảo luận video TikTok/YouTube"],
        [102, "Project_Marketing", "Chiến Dịch Marketing & Ads", "Nguyễn Thị B", "Topic chạy quảng cáo, content"],
        [103, "Project_Dev", "Phát Triển Hệ Thống / Tool", "Lê Văn C", "Topic kỹ thuật, website, bot"]
    ]

    for r_idx, row_data in enumerate(sample_mappings, 5):
        for c_idx, val in enumerate(row_data, 1):
            cell = ws_map.cell(row=r_idx, column=c_idx)
            cell.value = val
            cell.font = data_font
            cell.border = thin_border
            if c_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Column widths for mapping
    for col in ws_map.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_map.column_dimensions[col_letter].width = max(max_len + 4, 15)
    ws_map.row_dimensions[4].height = 28

    # ----------------------------------------------------
    # Helper to build standard Project sheets
    # ----------------------------------------------------
    project_sheets = [
        ("Project_Video", "DỰ ÁN: SẢN XUẤT VIDEO & TIKTOK", primary_header_fill, [
            ["TASK-001", "2026-08-27 09:30", "Làm video thẻ tín dụng", "Hoàn thiện 1 video ngắn về tính năng hoàn tiền thẻ tín dụng", "Sếp Vương", "Nguyễn Văn A", "2026-08-28 17:00", "Cao", "Chưa thực hiện", ""],
            ["TASK-002", "2026-08-27 10:15", "Edit video review tài khoản", "Cắt ghép và thêm subtitle cho video demo", "Sếp Vương", "Lê Văn B", "2026-08-29 12:00", "Trung bình", "Đang thực hiện", "Link source drive..."]
        ]),
        ("Project_Marketing", "DỰ ÁN: MARKETING & CONTENT", green_header_fill, [
            ["TASK-001", "2026-08-27 08:45", "Viết bài post khuyến mãi cuối tuần", "Chuẩn bị content bài đăng Facebook & Zalo", "Sếp Vương", "Trần Thị C", "2026-08-28 10:00", "Trung bình", "Đang thực hiện", ""]
        ]),
        ("Project_Dev", "DỰ ÁN: KỸ THUẬT & HỆ THỐNG", purple_header_fill, [
            ["TASK-001", "2026-08-27 08:00", "Kiểm tra webhook Telegram", "Cấu hình n8n webhook nhận event topic", "Sếp Vương", "Kỹ thuật D", "2026-08-27 18:00", "Khẩn cấp", "Hoàn thành", "Đã test OK"]
        ])
    ]

    task_headers = [
        "Mã Task", 
        "Thời Gian Tạo", 
        "Tên Task (AI Trích xuất)", 
        "Chi Tiết Yêu Cầu", 
        "Người Giao", 
        "Người Nhận (Assignee)", 
        "Hạn Chót (Deadline)", 
        "Mức Độ Ưu Tiên", 
        "Trạng Thái", 
        "Ghi Chú / Link Báo Cáo"
    ]

    # Priority & Status Validations
    dv_priority = DataValidation(type="list", formula1='"Thấp,Trung bình,Cao,Khẩn cấp"', allow_blank=True)
    dv_status = DataValidation(type="list", formula1='"Chưa thực hiện,Đang thực hiện,Chờ duyệt,Hoàn thành,Đã hủy"', allow_blank=True)

    for tab_title, proj_name, header_fill, sample_rows in project_sheets:
        ws = wb.create_sheet(title=tab_title)
        ws.views.sheetView[0].showGridLines = True
        ws.add_data_validation(dv_priority)
        ws.add_data_validation(dv_status)

        # Title
        ws["A1"] = proj_name
        ws["A1"].font = title_font
        ws["A2"] = f"Bảng quản lý công việc tự động cập nhật từ Telegram Topic '{tab_title}' qua n8n."
        ws["A2"].font = subtitle_font

        # Header
        for col_num, header in enumerate(task_headers, 1):
            cell = ws.cell(row=4, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.row_dimensions[4].height = 30

        # Sample rows
        for r_idx, row_data in enumerate(sample_rows, 5):
            for c_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = val
                cell.font = data_font
                cell.border = thin_border
                
                # Alignments
                if c_idx in [1, 2, 7, 8, 9]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        # Apply dropdowns to rows 5 to 500
        dv_priority.add(f"H5:H500")
        dv_status.add(f"I5:I500")

        # Column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            if col_letter in ['C', 'D', 'J']:
                ws.column_dimensions[col_letter].width = max(max_len + 4, 30)
            elif col_letter in ['B', 'G']:
                ws.column_dimensions[col_letter].width = 20
            else:
                ws.column_dimensions[col_letter].width = max(max_len + 4, 16)

    # Save workbook
    output_path = r"c:\Users\vuong\Downloads\ABF\Telegram_Project_Tasks_Template.xlsx"
    wb.save(output_path)
    print(f"Template created successfully at: {output_path}")

if __name__ == "__main__":
    create_project_task_template()
